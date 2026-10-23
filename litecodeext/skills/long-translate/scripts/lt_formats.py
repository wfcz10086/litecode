"""long-translate 格式层: 各格式 → 段列表 (extract), 译文 → 原格式输出 (write_output)。

txt / md / srt / pdf 用统一的 "pre + 正文 + post" 模型: 段里存前缀 (如 "## "、"1. "、字幕序号+时间轴)
和后缀 (换行分隔), 回写时 ''.join(pre + 译文 + post), 结构原样还原。
docx: 每个段落 (含表格单元格、页眉页脚) 一段, 译文写进第一个文字节点, 段落样式/首个 run 格式保留。
epub: 逐章 (spine 顺序) 解析 XHTML, 叶子块元素 (p/h1-6/li/td…) 一段; 行内标签 (<i>/<a>/脚注) 换成占位符
      让模型原样带回, 失败时退回纯文本。其他文件 (图片/CSS/目录) 原样拷贝。
pdf: 只取文字层 (pdftotext / PyMuPDF / pypdf), 重排成段落后按 Markdown 输出 (无法还原版式)。
"""
from __future__ import annotations

import html
import re
import shutil
import subprocess
import zipfile
from pathlib import Path

from lt_core import CHAPTER_RE, LETTER_RE

FORMATS = {".txt": "txt", ".text": "txt", ".md": "md", ".markdown": "md", ".srt": "srt",
           ".docx": "docx", ".epub": "epub", ".pdf": "pdf"}


def detect_format(path: str | Path) -> str:
    ext = Path(path).suffix.lower()
    if ext not in FORMATS:
        raise SystemExit(f"不支持的格式 {ext} (支持 {', '.join(sorted(FORMATS))})")
    return FORMATS[ext]


def output_ext(fmt: str) -> str:
    return {"pdf": ".md"}.get(fmt, "." + fmt)


def _seg(text, kind="p", pre="", post="", **kw):
    return dict(text=text, kind=kind, pre=pre, post=post, **kw)


# ── 占位符保护 ─────────────────────────────────────────────────
def protect(text: str, patterns: list[str]) -> tuple[str, dict]:
    """把匹配 patterns 的片段换成 ⟨n⟩, 返回 (新文本, {占位符: 原文})。"""
    ph: dict[str, str] = {}
    if not patterns:
        return text, ph
    rx = re.compile("|".join(f"(?:{p})" for p in patterns))

    def rep(m):
        k = f"⟨{len(ph)}⟩"
        ph[k] = m.group(0)
        return k
    return rx.sub(rep, text), ph


URL_PAT = r"https?://[^\s)\]>\"']*[^\s)\]>\"'.,;:!?。，；：！？]"   # 结尾标点不算 URL
MD_INLINE_PATS = [r"`[^`\n]+`", r"\]\([^)\s]*(?:\s+\"[^\"]*\")?\)", r"<https?://[^>]+>", r"<[^>\n]+>", URL_PAT,
                  r"\{#[^}]+\}"]


# ── txt ────────────────────────────────────────────────────────
def _reflow(block: str, cjk: bool) -> str:
    """硬换行的段落拼成一行; 短行 (诗歌/列表) 保留换行。"""
    lines = [l.strip() for l in block.splitlines()]
    if len(lines) > 1 and max(len(l) for l in lines) < 45:
        return "\n".join(lines)
    out = ""
    for l in lines:
        if not out:
            out = l
        elif out.endswith("-") and not out.endswith("--") and l[:1].islower():
            out = out[:-1] + l
        else:
            out += ("" if cjk else " ") + l
    return out


def extract_plain(text: str, cjk: bool = False, protect_urls: bool = True) -> list[dict]:
    text = text.replace("\r\n", "\n").replace("\r", "\n").lstrip("\ufeff")
    body_all = text.lstrip()
    lead = text[:len(text) - len(body_all)]
    parts = re.split(r"(\n[ \t]*\n\s*)", body_all)
    segs: list[dict] = []
    for k in range(0, len(parts), 2):
        block = parts[k]
        sep = parts[k + 1] if k + 1 < len(parts) else ""
        tail = block[len(block.rstrip()):]
        if not block.strip():
            if segs:
                segs[-1]["post"] += block + sep
            continue
        body = _reflow(block.strip(), cjk)
        kind = "heading" if "\n" not in body and CHAPTER_RE.match(body) else "p"
        ph = {}
        if protect_urls:
            body, ph = protect(body, [URL_PAT])
        segs.append(_seg(body, kind, pre=lead if not segs else "", post=tail + sep, ph=ph))
    return segs


# ── markdown ───────────────────────────────────────────────────
_MD_LIST = re.compile(r"^(\s*(?:[-*+]|\d+[.)])\s+(?:\[[ xX]\]\s+)?)")
_MD_HEAD = re.compile(r"^(#{1,6}\s+)")
_MD_QUOTE = re.compile(r"^(\s*(?:>\s?)+)")
_MD_FENCE = re.compile(r"^\s*(```|~~~)")
_MD_TABLE_SEP = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_MD_REFDEF = re.compile(r"^\s*\[[^\]]+\]:\s+\S+")


def extract_md(text: str) -> list[dict]:
    text = text.replace("\r\n", "\n").lstrip("﻿")
    lines = text.split("\n")
    segs: list[dict] = []
    i = 0

    def raw(s):                                  # 原样保留的片段
        if segs and segs[-1]["kind"] == "skip":
            segs[-1]["text"] += s
        else:
            segs.append(_seg(s, "skip"))

    def para(prefix, body, kind="p", post="\n"):
        body2, ph = protect(body, MD_INLINE_PATS)
        if not LETTER_RE.search(re.sub(r"⟨\d+⟩", "", body2)):
            raw(prefix + body + post)
            return
        segs.append(_seg(body2, kind, pre=prefix, post=post, ph=ph))

    if lines and lines[0].strip() == "---":      # YAML front matter
        for j in range(1, len(lines)):
            if lines[j].strip() in ("---", "..."):
                raw("\n".join(lines[:j + 1]) + "\n")
                i = j + 1
                break
    while i < len(lines):
        line = lines[i]
        last = i == len(lines) - 1
        nl = "" if last else "\n"
        m = _MD_FENCE.match(line)
        if m:
            j = i + 1
            while j < len(lines) and not lines[j].strip().startswith(m.group(1)):
                j += 1
            block = "\n".join(lines[i:j + 1]) + ("" if j >= len(lines) - 1 else "\n")
            segs.append(_seg(block, "code"))
            i = j + 1
            continue
        if not line.strip() or _MD_REFDEF.match(line) or _MD_TABLE_SEP.match(line) \
                or re.match(r"^\s*(\*\s*){3,}$|^\s*(-\s*){3,}$|^\s*(_\s*){3,}$", line) \
                or (line.startswith("    ") and (not segs or segs[-1]["post"].endswith("\n\n") or segs[-1]["kind"] == "skip")):
            raw(line + nl)
            i += 1
            continue
        h = _MD_HEAD.match(line)
        if h:
            para(h.group(1), line[h.end():], "heading", nl)
            i += 1
            continue
        if line.lstrip().startswith("|"):              # 表格行: 竖线当占位符保护, 整行一段
            body, ph = protect(line, [r"\|"] + MD_INLINE_PATS)
            if LETTER_RE.search(re.sub(r"⟨\d+⟩", "", body)):
                segs.append(_seg(body, "p", post=nl, ph=ph))
            else:
                raw(line + nl)
            i += 1
            continue
        pm = _MD_LIST.match(line) or _MD_QUOTE.match(line)
        prefix = pm.group(1) if pm else ""
        buf = [line[len(prefix):]]
        j = i + 1
        while j < len(lines):                       # 段落续行 (直到空行或新的块起始)
            nxt = lines[j]
            if (not nxt.strip() or _MD_HEAD.match(nxt) or _MD_FENCE.match(nxt) or _MD_LIST.match(nxt)
                    or nxt.lstrip().startswith("|") or _MD_QUOTE.match(nxt) != _MD_QUOTE.match(line) and _MD_QUOTE.match(nxt)):
                break
            buf.append(nxt.strip() if not _MD_QUOTE.match(nxt) else nxt[_MD_QUOTE.match(nxt).end():])
            j += 1
        body = " ".join(b.strip() for b in buf) if len(buf) > 1 else buf[0]
        para(prefix, body, "p", "" if j >= len(lines) else "\n")
        i = j
    return segs


# ── srt ────────────────────────────────────────────────────────
def extract_srt(text: str) -> list[dict]:
    text = text.replace("\r\n", "\n").lstrip("﻿").strip("\n") + "\n"
    segs = []
    for block in re.split(r"\n\s*\n", text):
        lines = block.strip("\n").split("\n")
        if len(lines) >= 2 and "-->" in lines[1]:
            body = "\n".join(lines[2:])
            if LETTER_RE.search(body):
                segs.append(_seg(body, "p", pre="\n".join(lines[:2]) + "\n", post="\n\n"))
            else:
                segs.append(_seg(block.strip("\n") + "\n\n", "skip"))
        elif block.strip():
            segs.append(_seg(block.strip("\n") + "\n\n", "skip"))
    return segs


# ── pdf ────────────────────────────────────────────────────────
def pdf_text(path: Path) -> str:
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-enc", "UTF-8", str(path), "-"], capture_output=True, timeout=600)
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout.decode("utf-8", "replace")
    try:
        import fitz  # type: ignore
        with fitz.open(str(path)) as d:
            return "\f".join(p.get_text() for p in d)
    except ImportError:
        pass
    try:
        from pypdf import PdfReader  # type: ignore
        return "\f".join((p.extract_text() or "") for p in PdfReader(str(path)).pages)
    except ImportError:
        raise SystemExit("PDF 需要 pdftotext / PyMuPDF / pypdf 之一")


def _pdf_paragraphs(text: str) -> str:
    """pdftotext 常常段间没有空行: 按"行明显短于版心宽度 + 以句末标点结尾"判段尾, 补空行。"""
    lines = text.split("\n")
    nonblank = [l for l in lines if l.strip()]
    if not nonblank or sum(1 for l in lines if not l.strip()) > len(nonblank) / 15:
        return text                                         # 已经有段落空行
    width = sorted(len(l.rstrip()) for l in nonblank)[int(len(nonblank) * 0.9)]
    out = []
    for l in lines:
        out.append(l)
        t = l.rstrip()
        if t and (len(t) < 0.8 * width and re.search(r"[.!?。！？:：”\"’'\])】]$", t) or len(t) < 0.45 * width):
            out.append("")
    return "\n".join(out)


def extract_pdf(path: Path, cjk: bool) -> list[dict]:
    raw_t = pdf_text(path)
    if len(raw_t.strip()) < 50:
        raise SystemExit("PDF 没有文字层 (扫描件), 先用 pdf-ocr-pipeline 技能 OCR 成文本再翻译")
    # 去掉页眉页码行 (纯数字行), 跨页段落接上
    pages = [re.sub(r"(?m)^\s*\d{1,4}\s*$", "", p) for p in raw_t.split("\f")]
    text = ""
    for p in pages:
        p = p.strip("\n")
        if not p:
            continue
        if text and not re.search(r"[.!?。！？:：”\"]\s*$", text) and p[:1].islower():
            text += " " + p.lstrip()
        else:
            text += ("\n\n" if text else "") + p
    segs = extract_plain(_pdf_paragraphs(text), cjk)
    for s in segs:
        if s["kind"] == "heading":
            s["pre"] = s["pre"] + "## "
    return segs


# ── docx ───────────────────────────────────────────────────────
def _docx_paragraphs(doc):
    """正文 + 表格 (递归) + 页眉页脚, 顺序固定, extract 与 write 用同一顺序。"""
    def from_container(c):
        for p in c.paragraphs:
            yield p
        for t in getattr(c, "tables", []):
            for row in t.rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:            # 合并单元格会重复出现
                        continue
                    seen.add(id(cell._tc))
                    yield from from_container(cell)
    yield from from_container(doc)
    for sec in doc.sections:
        for hf in (sec.header, sec.footer):
            if not hf.is_linked_to_previous:
                yield from from_container(hf)


def _wt_nodes(p):
    from docx.oxml.ns import qn
    return [t for t in p._p.iter(qn("w:t"))]


def _run_groups(p) -> list[list]:
    """按 run 格式 (rPr) 把文字节点分组, 相邻同格式的并成一组: [[w:t, ...], ...]。"""
    from docx.oxml.ns import qn
    from lxml import etree
    groups, last = [], None
    for t in _wt_nodes(p):
        r = t.getparent()
        rpr = r.find(qn("w:rPr")) if r is not None else None
        key = (etree.tostring(rpr) if rpr is not None else b"", r.getparent().tag if r is not None else "")
        if groups and key == last:
            groups[-1].append(t)
        else:
            groups.append([t])
        last = key
    return [g for g in groups if "".join(x.text or "" for x in g)] or groups[:1]


def extract_docx(path: Path) -> list[dict]:
    """多种格式的 run (加粗/斜体/链接…) 之间插 ⟨n⟩ 分隔标记 (seg["runs"]), 回写时按标记把译文分回各 run。"""
    import docx  # type: ignore
    doc = docx.Document(str(path))
    segs = []
    for k, p in enumerate(_docx_paragraphs(doc)):
        groups = _run_groups(p)
        texts = ["".join(t.text or "" for t in g) for g in groups]
        txt = "".join(texts)
        style = (p.style.name if p.style is not None else "") or ""
        heading = bool(re.match(r"(Heading|Title|标题)", style, re.I))
        kind = ("heading" if heading else "p") if LETTER_RE.search(txt) else "skip"
        if kind == "skip" or len(groups) <= 1 or len(groups) > 12:
            body, ph = protect(txt, [URL_PAT])
            segs.append(_seg(body, kind, loc=k, ph=ph, style=style))
            continue
        ph, runs, parts = {}, [], []
        for g_i, gt in enumerate(texts):
            b, sub = protect(gt, [URL_PAT])
            for old in sorted(sub, key=lambda x: -len(x)):        # 每组各自编号 → 重排成全段唯一编号
                new = f"⟨{len(ph) + 100}⟩"
                b = b.replace(old, new)
                ph[new] = sub[old]
            if g_i:
                mk = f"⟨{len(runs) + 1}⟩"
                runs.append(mk)
                parts.append(mk)
            parts.append(b)
        segs.append(_seg("".join(parts), kind, loc=k, ph=ph, style=style, runs=runs))
    return segs


def _split_runs(text: str, runs: list[str]) -> list[str] | None:
    """按 ⟨1⟩⟨2⟩… 把译文切回各 run; 标记缺失/乱序 → None。"""
    pos, pieces = 0, []
    for mk in runs:
        j = text.find(mk, pos)
        if j < 0:
            return None
        pieces.append(text[pos:j])
        pos = j + len(mk)
    pieces.append(text[pos:])
    return pieces


def write_docx(src: Path, segs: list[dict], trans: dict, out: Path):
    import docx  # type: ignore
    from docx.oxml.ns import qn
    doc = docx.Document(str(src))
    by_loc = {s["loc"]: s for s in segs if not s.get("cont")}
    xml_space = "{http://www.w3.org/XML/1998/namespace}space"
    for k, p in enumerate(_docx_paragraphs(doc)):
        s = by_loc.get(k)
        if not s or s["i"] not in trans:
            continue
        nodes = _wt_nodes(p)
        if not nodes:
            continue
        text = trans[s["i"]]
        pieces = _split_runs(text, s.get("runs") or []) if s.get("runs") else None
        if pieces is not None:
            groups = _run_groups(p)
            if len(groups) == len(pieces):
                for g, piece in zip(groups, pieces):
                    for t in g:
                        t.text = ""
                    g[0].text = piece
                    g[0].set(xml_space, "preserve")
                continue
        text = re.sub(r"⟨\d+⟩", "", text)
        # 标记对不上: 整段写进不在超链接里的第一个文字节点, 保留其 run 格式
        tgt = next((t for t in nodes if t.getparent().getparent().tag != qn("w:hyperlink")), nodes[0])
        for t in nodes:
            t.text = ""
        tgt.text = text
        tgt.set(xml_space, "preserve")
    doc.save(str(out))


# ── epub ───────────────────────────────────────────────────────
XHTML_NS = "http://www.w3.org/1999/xhtml"
EPUB_BLOCKS = {"p", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "td", "th", "dt", "dd",
               "figcaption", "caption", "div", "pre", "title"}
_TAG_RE = re.compile(r"<[^>]+>")


def _local(tag) -> str:
    return tag.rsplit("}", 1)[-1].lower() if isinstance(tag, str) else ""


def epub_spine(z: zipfile.ZipFile) -> tuple[str, list[str]]:
    from lxml import etree
    cont = etree.fromstring(z.read("META-INF/container.xml"))
    opf = cont.xpath("//*[local-name()='rootfile']/@full-path")[0]
    root = etree.fromstring(z.read(opf))
    base = opf.rsplit("/", 1)[0] + "/" if "/" in opf else ""
    items = {i.get("id"): i for i in root.xpath("//*[local-name()='manifest']/*[local-name()='item']")}
    docs = []
    for ref in root.xpath("//*[local-name()='spine']/*[local-name()='itemref']/@idref"):
        it = items.get(ref)
        if it is not None and "html" in (it.get("media-type") or ""):
            docs.append(base + it.get("href").split("#")[0])
    return opf, docs


def _epub_leaf_blocks(root):
    out = []
    for el in root.iter():
        name = _local(el.tag)
        if name not in EPUB_BLOCKS or name == "title":
            continue
        if any(_local(d.tag) in EPUB_BLOCKS for d in el.iterdescendants()):
            continue                                       # 只要最内层块
        out.append(el)
    return out


def _inner_xml(el) -> str:
    from lxml import etree
    return html.escape(el.text or "", quote=False) + "".join(
        etree.tostring(c, encoding="unicode", with_tail=True) for c in el)


def _tagged_to_text(inner: str) -> tuple[str, dict]:
    """行内标签 → 占位符; 文本里的实体反转义, 给模型看干净文本。"""
    ph, out, pos = {}, [], 0
    for m in _TAG_RE.finditer(inner):
        out.append(html.unescape(inner[pos:m.start()]))
        k = f"⟨{len(ph)}⟩"
        ph[k] = m.group(0)
        out.append(k)
        pos = m.end()
    out.append(html.unescape(inner[pos:]))
    return "".join(out), ph


def extract_epub(path: Path) -> list[dict]:
    from lxml import etree
    segs = []
    with zipfile.ZipFile(path) as z:
        _, docs = epub_spine(z)
        for d_i, name in enumerate(docs):
            try:
                root = etree.fromstring(z.read(name))
            except etree.XMLSyntaxError:
                continue
            first = True
            for b_i, el in enumerate(_epub_leaf_blocks(root)):
                text, ph = _tagged_to_text(_inner_xml(el))
                if not LETTER_RE.search(re.sub(r"⟨\d+⟩", "", text)):
                    continue
                kind = "heading" if _local(el.tag) in ("h1", "h2", "h3") else ("code" if _local(el.tag) == "pre" else "p")
                segs.append(_seg(text.strip(), kind, loc=[name, b_i], ph=ph, brk=first))
                first = False
    return segs


def _set_inner(el, s: str):
    from lxml import etree
    ns = el.nsmap.get(None) or XHTML_NS
    wrap = etree.fromstring(f'<w xmlns="{ns}">{s}</w>'.encode("utf-8"))
    for c in list(el):
        el.remove(c)
    el.text = wrap.text
    for c in list(wrap):
        el.append(c)


def write_epub(src: Path, segs: list[dict], trans: dict, out: Path, tgt_lang: str = "", raw_trans: dict | None = None):
    """trans: 已还原占位符的 XHTML 片段 (文本已转义); 解析失败则退回纯文本。"""
    from lxml import etree
    by_doc: dict[str, dict[int, dict]] = {}
    for s in segs:
        if not s.get("cont") and s["i"] in trans:
            by_doc.setdefault(s["loc"][0], {})[s["loc"][1]] = s
    with zipfile.ZipFile(src) as z, zipfile.ZipFile(out, "w") as zo:
        opf, _ = epub_spine(z)
        names = z.namelist()
        if "mimetype" in names:
            zo.writestr(zipfile.ZipInfo("mimetype"), z.read("mimetype"), compress_type=zipfile.ZIP_STORED)
        for name in names:
            if name == "mimetype":
                continue
            data = z.read(name)
            if name in by_doc:
                root = etree.fromstring(data)
                blocks = _epub_leaf_blocks(root)
                for b_i, s in by_doc[name].items():
                    el = blocks[b_i]
                    try:
                        _set_inner(el, trans[s["i"]])
                    except etree.XMLSyntaxError:
                        plain = (raw_trans or {}).get(s["i"]) or _TAG_RE.sub("", trans[s["i"]])
                        _set_inner(el, html.escape(re.sub(r"⟨\d+⟩", "", plain), quote=False))
                data = etree.tostring(root.getroottree(), xml_declaration=True, encoding="utf-8")
            elif name == opf and tgt_lang:
                data = re.sub(rb"(<dc:language[^>]*>)[^<]*(</dc:language>)",
                              b"\\g<1>" + tgt_lang.encode() + b"\\g<2>", data)
            zo.writestr(name, data, compress_type=zipfile.ZIP_DEFLATED)


# ── 统一入口 ───────────────────────────────────────────────────
def extract(path: str | Path, cjk_src: bool | None = None) -> tuple[str, list[dict]]:
    path = Path(path)
    fmt = detect_format(path)
    if fmt in ("txt", "md", "srt"):
        text = path.read_text(encoding="utf-8", errors="replace")
        if cjk_src is None:
            from lt_core import detect_lang
            cjk_src = detect_lang(text) != "en"
        segs = {"txt": lambda: extract_plain(text, cjk_src), "md": lambda: extract_md(text),
                "srt": lambda: extract_srt(text)}[fmt]()
    elif fmt == "pdf":
        segs = extract_pdf(path, bool(cjk_src))
    elif fmt == "docx":
        segs = extract_docx(path)
    else:
        segs = extract_epub(path)
    for i, s in enumerate(segs):
        s["i"] = i
    return fmt, segs


def merge_continuations(segs: list[dict], trans: dict, cjk_tgt: bool) -> dict:
    """超长段被拆成 head+cont 段; 把 cont 的译文接回 head。返回 {head_i: 译文}。"""
    out, head = {}, None
    for s in segs:
        t = trans.get(s["i"])
        if s.get("cont") and head is not None:
            if t is not None and head in out:
                out[head] = out[head] + ("" if cjk_tgt else " ") + t
            continue
        head = s["i"]
        if t is not None:
            out[head] = t
    return out


def rebuild_text(segs: list[dict], final: dict) -> str:
    """pre + 正文 + post 拼回整篇 (txt/md/srt/pdf)。
    cont 段 (超长段拆出的后半) 的译文已由 merge_continuations 并入 head; 没译文时回落原文。"""
    out, head = [], None
    for s in segs:
        if s.get("cont") and head is not None:
            if head["i"] not in final:               # head 没有译文 → 补上本段原文 (用 head 的占位符表)
                out.append(_restore_src(dict(s, ph=head.get("ph"))))
            out.append(s.get("post", ""))
            continue
        head = s
        body = final[s["i"]] if s["i"] in final else _restore_src(s)
        out.append(s.get("pre", "") + body + s.get("post", ""))
    return "".join(out)


def _restore_src(s: dict) -> str:
    t = s["text"]
    for k, v in (s.get("ph") or {}).items():
        t = t.replace(k, v)
    return t


def write_output(fmt: str, src: Path, segs: list[dict], final: dict, out: Path, tgt_lang: str = ""):
    out.parent.mkdir(parents=True, exist_ok=True)
    if fmt == "docx":
        write_docx(src, segs, final, out)
    elif fmt == "epub":
        write_epub(src, segs, final, out, tgt_lang)
    else:
        out.write_text(rebuild_text(segs, final), encoding="utf-8")
