#!/usr/bin/env python3
"""html2docx.py — 文档 HTML → DOCX (标题层级 / 段落 / 列表 / 表格 / 图片 / 图注 / 页脚页码)。

引擎:
  --engine auto    (默认) 有 pandoc 用 pandoc + 参考模板 (字体/标题色取自页面 --ds-* tokens), 否则 python-docx 映射
  --engine pandoc  强制 pandoc            --engine python  强制 python-docx 映射 (版式更贴近: 栅格→无框表格)
用法:
  python3 html2docx.py doc.html -o doc.docx
  python3 html2docx.py doc.html -o doc.docx --engine python --page-breaks --font "Microsoft YaHei"
约定: data-docx="skip" 的元素不导出; 演示用的图标 SVG、aside.notes、script/style 自动丢弃;
      .grid-2/.grid-3 或 data-docx="columns" 的栅格 → 无框表格并排。
输出: stdout 打印 JSON {"docx", "engine", "headings", "paragraphs", "tables", "images"}。
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
import urllib.parse
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _ds  # noqa: E402
from _img_make import css_tokens  # noqa: E402

BLOCK_TAGS = {"p", "div", "section", "article", "header", "footer", "main", "aside", "nav", "figure",
              "figcaption", "ul", "ol", "li", "table", "blockquote", "pre", "h1", "h2", "h3", "h4", "h5", "h6",
              "hr", "img", "dl", "dt", "dd", "svg", "canvas", "video"}
INLINE_FMT = {"b": "bold", "strong": "bold", "i": "italic", "em": "italic", "u": "underline",
              "code": "code", "kbd": "code", "sup": "sup", "sub": "sub", "mark": "bold", "small": "small"}


def clean_soup(html_path: Path):
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html_path.read_text(encoding="utf-8"), "html.parser")
    for sel in ("script", "style", "noscript", "template", "aside.notes", "[data-docx=skip]",
                "svg.ds-icon", "svg[data-icon]", "[data-pptx-only]", ".page-footer", ".slide .footer"):
        for el in soup.select(sel):
            el.decompose()
    return soup


def _hex(s: str | None, default: str) -> str:
    m = re.match(r"\s*#?([0-9a-fA-F]{6})\b", s or "")
    return (m.group(1) if m else default).upper()


def _img_path(src: str, base: Path) -> Path | None:
    if not src or src.startswith(("http://", "https://", "data:")):
        return None
    if src.startswith("file://"):
        return Path(urllib.parse.unquote(urllib.parse.urlparse(src).path))
    p = (base / urllib.parse.unquote(src)).resolve()
    return p if p.exists() else None


def grid_classes(css: str) -> dict:
    """<style> 里 display:grid 的类 → 列数 (repeat(N, …) 或 1fr 1fr 计数)。"""
    out = {}
    for m in re.finditer(r"\.([\w-]+)\s*\{([^}]*)\}", css):
        body = m.group(2)
        g = re.search(r"grid-template-columns\s*:\s*([^;]+)", body)
        if not g:
            continue
        r = re.search(r"repeat\(\s*(\d+)", g.group(1))
        n = int(r.group(1)) if r else len(g.group(1).split())
        if 1 < n <= 6:
            out[m.group(1)] = n
    return out


# ── python-docx 映射 ───────────────────────────────────
class DocxWriter:
    def __init__(self, base: Path, tokens: dict, font: str, page_breaks: bool):
        from docx import Document
        from docx.shared import Mm, Pt, RGBColor
        self.base, self.font, self.page_breaks = base, font, page_breaks
        self.primary = RGBColor.from_string(_hex(tokens.get("primary"), "1F2A44"))
        self.accent = RGBColor.from_string(_hex(tokens.get("accent"), "4F7EF7"))
        self.text_c = RGBColor.from_string(_hex(tokens.get("text"), "1A1A1A"))
        self.muted = RGBColor.from_string(_hex(tokens.get("text-3"), "999999"))
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.left_margin = sec.right_margin = Mm(20)
        sec.top_margin, sec.bottom_margin = Mm(20), Mm(18)
        self.text_width_mm = 170
        self.stats = {"headings": 0, "paragraphs": 0, "tables": 0, "images": 0}
        self.grid_classes: dict = {}
        st = self.doc.styles
        self._style_font(st["Normal"], Pt(10.5), self.text_c)
        st["Normal"].paragraph_format.space_after = Pt(6)
        st["Normal"].paragraph_format.line_spacing = 1.35
        for i, size in ((1, 20), (2, 15), (3, 12.5), (4, 11.5)):
            self._style_font(st[f"Heading {i}"], Pt(size), self.primary, bold=True)
            st[f"Heading {i}"].paragraph_format.space_before = Pt(14 if i <= 2 else 10)
            st[f"Heading {i}"].paragraph_format.space_after = Pt(6)
        self._style_font(st["Title"], Pt(28), self.primary, bold=True)
        self._style_font(st["Caption"], Pt(9), self.muted, bold=False)
        self._footer()

    def _style_font(self, style, size, color, bold=None):
        from docx.oxml.ns import qn
        style.font.name = self.font
        style.font.size = size
        style.font.color.rgb = color
        if bold is not None:
            style.font.bold = bold
        rpr = style.element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is None:
            rf = rpr.makeelement(qn("w:rFonts"), {})
            rpr.insert(0, rf)
        for k in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(k), self.font)
        for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rf.attrib.pop(qn(k), None)

    def _footer(self):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), "PAGE")
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = "1"
        r.append(t)
        fld.append(r)
        p._p.append(fld)

    # 行内
    def runs(self, par, node, fmt: dict | None = None):
        from bs4 import NavigableString, Tag
        from docx.shared import Pt, RGBColor
        fmt = dict(fmt or {})
        if isinstance(node, NavigableString):
            txt = re.sub(r"\s+", " ", str(node))
            if txt.strip() or (txt and par.runs):
                r = par.add_run(txt)
                r.bold = fmt.get("bold") or None
                r.italic = fmt.get("italic") or None
                r.underline = fmt.get("underline") or None
                if fmt.get("code"):
                    r.font.name = "Consolas"
                if fmt.get("sup"):
                    r.font.superscript = True
                if fmt.get("sub"):
                    r.font.subscript = True
                if fmt.get("small"):
                    r.font.size = Pt(8.5)
                if fmt.get("color"):
                    r.font.color.rgb = RGBColor.from_string(fmt["color"])
                if fmt.get("link"):
                    r.font.color.rgb = self.accent
                    r.underline = True
            return
        if not isinstance(node, Tag):
            return
        if node.name == "br":
            par.add_run().add_break()
            return
        if node.name in ("img", "svg"):
            return
        if node.name in INLINE_FMT:
            fmt[INLINE_FMT[node.name]] = True
        if node.name == "a" and node.get("href", "").startswith("http"):
            fmt["link"] = True
        m = re.search(r"(?:^|;)\s*color\s*:\s*#([0-9a-fA-F]{6})", node.get("style") or "")
        if m:
            fmt["color"] = m.group(1).upper()
        for c in node.children:
            self.runs(par, c, fmt)

    def _is_inline_only(self, el) -> bool:
        from bs4 import Tag
        return not any(isinstance(c, Tag) and (c.name in BLOCK_TAGS or self._bg_image(c)) for c in el.descendants)

    @staticmethod
    def _bg_image(el) -> str | None:
        m = re.search(r"background-image\s*:\s*url\(['\"]?([^'\")]+)", el.get("style") or "")
        return m.group(1) if m else None

    def _strip_trailing(self, par):
        if par.runs and par.runs[-1].text:
            par.runs[-1].text = par.runs[-1].text.rstrip()
        if par.runs and par.runs[0].text:
            par.runs[0].text = par.runs[0].text.lstrip()

    def paragraph(self, box, el, style: str | None = None, align=None):
        par = box.add_paragraph(style=style) if style else box.add_paragraph()
        for c in el.children:
            self.runs(par, c)
        self._strip_trailing(par)
        if not "".join(r.text for r in par.runs).strip() and not par.runs:
            par._p.getparent().remove(par._p)
            return None
        if align is not None:
            par.alignment = align
        self.stats["paragraphs"] += 1
        return par

    def picture(self, box, path: Path, width_mm: float, caption: str = ""):
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Mm
        from PIL import Image
        try:
            with Image.open(path) as im:
                w, h = im.size
        except Exception:
            return
        width_mm = min(width_mm, self.text_width_mm)
        if h > w * 1.25:                               # 竖图别太大
            width_mm = min(width_mm, 80)
        par = box.add_paragraph()
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.add_run().add_picture(str(path), width=Mm(width_mm))
        self.stats["images"] += 1
        if caption:
            cp = box.add_paragraph(caption, style="Caption")
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 块级
    def block(self, box, el, width_mm: float, list_level: int = 0):
        from bs4 import NavigableString, Tag
        if isinstance(el, NavigableString):
            if str(el).strip():
                par = box.add_paragraph(re.sub(r"\s+", " ", str(el)).strip())
                self.stats["paragraphs"] += 1
                return par
            return None
        if not isinstance(el, Tag):
            return None
        name = el.name
        classes = el.get("class") or []
        if name in ("svg", "canvas", "video", "hr", "script", "style"):
            return None
        if name in ("h1", "h2", "h3", "h4", "h5", "h6"):
            lvl = min(4, int(name[1]))
            par = box.add_paragraph(style=f"Heading {lvl}")
            for c in el.children:
                self.runs(par, c)
            self._strip_trailing(par)
            self.stats["headings"] += 1
            return par
        if name == "img":
            p = _img_path(el.get("src", ""), self.base)
            if p:
                self.picture(box, p, width_mm)
            return None
        bg = self._bg_image(el)
        if bg and _img_path(bg, self.base):
            self.picture(box, _img_path(bg, self.base), width_mm)
        if name == "figure":
            imgs = el.find_all("img")
            cap = el.find("figcaption")
            cap_txt = cap.get_text(" ", strip=True) if cap else ""
            credit = el.select_one("[data-credit]")
            if credit and credit.get_text(strip=True) and credit.get_text(strip=True) not in cap_txt:
                cap_txt = (cap_txt + "  " if cap_txt else "") + credit.get_text(" ", strip=True)
            for i, im in enumerate(imgs):
                p = _img_path(im.get("src", ""), self.base)
                if p:
                    self.picture(box, p, width_mm / max(1, len(imgs)) if len(imgs) > 1 else width_mm,
                                 cap_txt if i == len(imgs) - 1 else "")
            return None
        if name in ("ul", "ol"):
            for li in el.find_all("li", recursive=False):
                style = ("List Number" if name == "ol" else "List Bullet") + (f" {list_level + 1}" if list_level else "")
                par = box.add_paragraph(style=style)
                for c in li.children:
                    if isinstance(c, Tag) and c.name in ("ul", "ol"):
                        continue
                    if isinstance(c, Tag) and c.name in BLOCK_TAGS and c.name not in ("ul", "ol"):
                        for cc in c.children:
                            self.runs(par, cc)
                        continue
                    self.runs(par, c)
                self._strip_trailing(par)
                self.stats["paragraphs"] += 1
                for sub in li.find_all(["ul", "ol"], recursive=False):
                    self.block(box, sub, width_mm, min(2, list_level + 1))
            return None
        if name == "table":
            self.table(box, el)
            return None
        if name == "blockquote" and self._is_inline_only(el):
            return self.paragraph(box, el, "Intense Quote")
        if name == "pre":
            par = box.add_paragraph()
            r = par.add_run(el.get_text())
            r.font.name = "Consolas"
            self.stats["paragraphs"] += 1
            return par
        if name == "section" and "page" in classes and self.page_breaks and self.stats["paragraphs"] > 0:
            from docx.enum.text import WD_BREAK
            box.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        ncol = self._grid_cols(el)
        if ncol:
            self.columns(box, el, ncol, width_mm)
            return None
        if self._is_inline_only(el):
            if not el.get_text(strip=True):
                return None
            style = "Caption" if name == "figcaption" or {"credit", "caption", "credits"} & set(classes) else None
            return self.paragraph(box, el, style)
        for c in el.children:
            self.block(box, c, width_mm, list_level)
        return None

    def _grid_cols(self, el) -> int:
        from bs4 import Tag
        classes = el.get("class") or []
        kids = [c for c in el.children if isinstance(c, Tag)]
        if len(kids) < 2:
            return 0
        for c in classes:
            m = re.match(r"grid-(\d)$", c)
            if m:
                return min(int(m.group(1)), len(kids))
            if c in self.grid_classes:                  # 页面 <style> 里定义成 grid 的类
                return min(self.grid_classes[c], len(kids))
        if el.get("data-docx") == "columns" or re.search(r"display\s*:\s*grid", el.get("style") or ""):
            return min(4, len(kids))
        return 0

    def columns(self, box, el, ncol: int, width_mm: float):
        from bs4 import Tag
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        kids = [c for c in el.children if isinstance(c, Tag)]
        nrow = (len(kids) + ncol - 1) // ncol
        t = box.add_table(rows=nrow, cols=ncol)
        tblPr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            b.set(qn("w:val"), "nil")
            borders.append(b)
        tblPr.append(borders)
        self.stats["tables"] += 1
        for i, k in enumerate(kids):
            cell = t.cell(i // ncol, i % ncol)
            first = cell.paragraphs[0]
            self.block(cell, k, width_mm / ncol - 4)
            if not first.runs and len(cell.paragraphs) > 1:
                first._p.getparent().remove(first._p)

    def table(self, box, el):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        rows = el.find_all("tr")
        if not rows:
            return
        ncol = max(sum(int(c.get("colspan", 1) or 1) for c in r.find_all(["td", "th"], recursive=False)) for r in rows)
        t = box.add_table(rows=len(rows), cols=ncol)
        t.style = "Table Grid"
        borders = OxmlElement("w:tblBorders")             # 细浅灰框线, 不要默认的黑粗线
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b = OxmlElement(f"w:{side}")
            for k, v in (("w:val", "single"), ("w:sz", "4"), ("w:color", "D9D9D9")):
                b.set(qn(k), v)
            borders.append(b)
        t._tbl.tblPr.append(borders)
        self.stats["tables"] += 1
        if rows[0].find("th"):                             # 表头跨页重复
            trPr = t.rows[0]._tr.get_or_add_trPr()
            trPr.append(OxmlElement("w:tblHeader"))
        taken: set = set()
        for i, r in enumerate(rows):
            j = 0
            for c in r.find_all(["td", "th"], recursive=False):
                while (i, j) in taken:
                    j += 1
                if j >= ncol:
                    break
                cs, rs = int(c.get("colspan", 1) or 1), int(c.get("rowspan", 1) or 1)
                cell = t.cell(i, j)
                end = t.cell(min(len(rows) - 1, i + rs - 1), min(ncol - 1, j + cs - 1))
                if end is not cell and (rs > 1 or cs > 1):
                    cell = cell.merge(end)
                for di in range(rs):
                    for dj in range(cs):
                        taken.add((i + di, j + dj))
                par = cell.paragraphs[0]
                for cc in c.children:
                    self.runs(par, cc, {"bold": c.name == "th"})
                self._strip_trailing(par)
                if c.name == "th":
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), str(self.primary))
                    cell._tc.get_or_add_tcPr().append(shd)
                    for run in par.runs:
                        from docx.shared import RGBColor
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                j += cs

    def write(self, soup, out: Path):
        body = soup.body or soup
        for c in body.children:
            self.block(self.doc, c, self.text_width_mm)
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))
        return self.stats


# ── pandoc ───────────────────────────────────────────
def pandoc_reference(tmp: Path, tokens: dict, font: str) -> Path | None:
    ref = tmp / "reference.docx"
    try:
        subprocess.run(["pandoc", "-o", str(ref), "--print-default-data-file", "reference.docx"],
                       capture_output=True, timeout=30, check=True)
        from docx import Document
        from docx.oxml.ns import qn
        from docx.shared import RGBColor
        d = Document(str(ref))
        prim = RGBColor.from_string(_hex(tokens.get("primary"), "1F2A44"))
        for s in d.styles:
            try:
                if s.type != 1:                            # 只改段落样式
                    continue
            except Exception:
                continue
            rpr = s.element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = rpr.makeelement(qn("w:rFonts"), {})
                rpr.insert(0, rf)
            for k in ("w:ascii", "w:hAnsi", "w:eastAsia"):
                rf.set(qn(k), font)
            for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme"):
                rf.attrib.pop(qn(k), None)
            if s.name.startswith("Heading") or s.name == "Title":
                s.font.color.rgb = prim
        d.save(str(ref))
        return ref
    except Exception as e:
        _ds.eprint(f"[html2docx] 生成参考模板失败, 用 pandoc 默认: {e}")
        return None


def via_pandoc(soup, base: Path, out: Path, tokens: dict, font: str) -> dict:
    tmp = Path(tempfile.mkdtemp(prefix="html2docx_"))
    try:
        for div in soup.select("[style*=background-image]"):     # 背景图 → <img>, pandoc 才认
            m = re.search(r"url\(['\"]?([^'\")]+)", div.get("style", ""))
            if m:
                img = soup.new_tag("img", src=m.group(1))
                div.insert(0, img)
        src = tmp / "in.html"
        src.write_text(str(soup), encoding="utf-8")
        cmd = ["pandoc", str(src), "-f", "html", "-t", "docx", "-o", str(out), f"--resource-path={base}"]
        ref = pandoc_reference(tmp, tokens, font)
        if ref:
            cmd.append(f"--reference-doc={ref}")
        out.parent.mkdir(parents=True, exist_ok=True)
        subprocess.run(cmd, check=True, capture_output=True, timeout=180)
        from docx import Document
        d = Document(str(out))
        return {"headings": sum(1 for p in d.paragraphs if p.style.name.startswith("Heading")),
                "paragraphs": len(d.paragraphs), "tables": len(d.tables), "images": len(d.inline_shapes)}
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def convert(html: str, out: str, engine: str = "auto", font: str = "Microsoft YaHei",
            page_breaks: bool = False) -> dict:
    src = Path(html).resolve()
    tokens = css_tokens(src.read_text(encoding="utf-8"))
    soup = clean_soup(src)
    use = engine
    if engine == "auto":
        use = "pandoc" if shutil.which("pandoc") else "python"
    o = Path(out).resolve()
    if use == "pandoc":
        stats = via_pandoc(soup, src.parent, o, tokens, font)
    else:
        w = DocxWriter(src.parent, tokens, font, page_breaks)
        raw = src.read_text(encoding="utf-8")            # clean_soup 已删 <style>, 栅格类从原文取
        w.grid_classes = grid_classes(" ".join(re.findall(r"<style[^>]*>(.*?)</style>", raw, re.S)))
        stats = w.write(soup, o)
    return {"docx": str(o), "engine": use, **stats}


def main(argv=None):
    ap = argparse.ArgumentParser(description="doc-studio: 文档 HTML → DOCX")
    ap.add_argument("html")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--engine", default="auto", choices=["auto", "pandoc", "python"])
    ap.add_argument("--font", default="Microsoft YaHei")
    ap.add_argument("--page-breaks", action="store_true", help="每个 section.page 前分页 (python 引擎)")
    a = ap.parse_args(argv)
    print(_ds.dump_json(convert(a.html, a.out, a.engine, a.font, a.page_breaks)))


if __name__ == "__main__":
    main()
