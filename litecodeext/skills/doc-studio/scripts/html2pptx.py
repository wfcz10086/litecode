#!/usr/bin/env python3
"""html2pptx.py — HTML 幻灯片 (section.slide) → PPTX。

  --mode editable  按浏览器里每个元素的计算样式与坐标, 用 python-pptx 还原: 文字框 (字体/字号/颜色/加粗/对齐/行距)、
                   图片 (object-fit 裁剪/圆角)、色块与边框、原生表格、背景; 渐变/滤镜/SVG/图表等单独截图成图片
  --mode image     高保真: 整页截图做背景, 全部文字放备注 + 图片下方的隐藏文本框 (可搜索、可复制)
  --mode auto      (默认) 逐页判断: 元素过多 / 大面积只能截图 → 该页用 image, 否则 editable
用法:
  python3 html2pptx.py deck.html -o deck.pptx
  python3 html2pptx.py deck.html -o deck.pptx --mode image
  python3 html2pptx.py deck.html -o deck.pptx --font "Microsoft YaHei" --keep-fonts
输出: stdout 打印 JSON {"pptx", "slides", "modes": [...], "reasons": [...]}。
备注: <aside class="notes">…</aside> 写在 section 里会进 PPT 演讲者备注 (自身别显示: display:none)。
"""
from __future__ import annotations

import argparse
import base64
import shutil
import sys
import tempfile
import urllib.parse
import urllib.request
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import _ds  # noqa: E402

EXTRACT_JS = (Path(__file__).resolve().parent / "_pptx_extract.js").read_text(encoding="utf-8")
EMU_PER_PX = 9525                                  # 1 CSS px = 1/96 in = 9525 EMU
SAFE_LATIN = {"arial", "helvetica", "georgia", "times new roman", "segoe ui", "calibri", "verdana",
              "consolas", "courier new", "tahoma", "microsoft yahei", "simsun", "simhei"}
CAP_CSS = """
html.ds-cap, html.ds-cap body { background: transparent !important; }
html.ds-cap body * { visibility: hidden !important; transition: none !important; animation: none !important; }
html.ds-cap [data-ds-capnow] { visibility: visible !important; }
html.ds-cap [data-ds-capnow][data-ds-mode=self] { color: transparent !important; text-shadow: none !important; }
html.ds-cap [data-ds-capnow][data-ds-mode=full] * { visibility: visible !important; }
"""


# ── 浏览器阶段: 抽取 + 截图 ───────────────────────────────
def extract(page, selector: str) -> list[dict]:
    return page.evaluate(EXTRACT_JS, {"selector": selector})


def capture_el(page, cap_id: int, mode: str, dest: Path) -> Path | None:
    r = page.evaluate("""([id, mode]) => {
        const e = document.querySelector(`[data-ds-cap="${id}"]`);
        if (!e) return null;
        e.setAttribute('data-ds-capnow', '1'); e.setAttribute('data-ds-mode', mode);
        document.documentElement.classList.add('ds-cap');
        const r = e.getBoundingClientRect();
        return {x: r.left + scrollX, y: r.top + scrollY, w: r.width, h: r.height};
    }""", [cap_id, mode])
    try:
        if not r or r["w"] < 1 or r["h"] < 1:
            return None
        page.screenshot(path=str(dest), full_page=True, omit_background=True,
                        clip={"x": r["x"], "y": r["y"], "width": r["w"], "height": r["h"]})
        return dest
    finally:
        page.evaluate("""(id) => {
            const e = document.querySelector(`[data-ds-cap="${id}"]`);
            if (e) { e.removeAttribute('data-ds-capnow'); e.removeAttribute('data-ds-mode'); }
            document.documentElement.classList.remove('ds-cap');
        }""", cap_id)


def shoot_section(page, selector: str, i: int, dest: Path) -> Path:
    r = page.evaluate("""([sel, i]) => { const s = document.querySelectorAll(sel)[i];
        const r = s.getBoundingClientRect(); return {x: r.left + scrollX, y: r.top + scrollY, w: r.width, h: r.height}; }""",
                      [selector, i])
    page.screenshot(path=str(dest), full_page=True, clip={"x": r["x"], "y": r["y"], "width": r["w"], "height": r["h"]})
    return dest


def choose_mode(s: dict, mode: str) -> tuple[str, str]:
    if mode != "auto":
        return mode, "指定"
    if s.get("error"):
        return "image", f"抽取出错: {s['error'][:60]}"
    m = s.get("metrics") or {}
    if m.get("items", 0) > 150:
        return "image", f"元素过多 ({m['items']})"
    area = (s.get("w") or 1) * (s.get("h") or 1)
    if m.get("capFullArea", 0) > 0.5 * area:
        return "image", "大面积只能截图 (SVG/滤镜/渐变字)"
    return "editable", f"{m.get('textBlocks', 0)} 个文字块, {m.get('caps', 0)} 处截图"


# ── 构建 PPTX ───────────────────────────────────────
class Builder:
    def __init__(self, w_px: float, h_px: float, font: str, keep_fonts: bool, tmp: Path):
        from pptx import Presentation
        self.prs = Presentation()
        self.prs.slide_width = int(w_px * EMU_PER_PX)
        self.prs.slide_height = int(h_px * EMU_PER_PX)
        self.font, self.keep_fonts, self.tmp = font, keep_fonts, tmp
        self._dl: dict = {}

    @staticmethod
    def E(v: float) -> int:
        return int(round(v * EMU_PER_PX))

    def new_slide(self):
        return self.prs.slides.add_slide(self.prs.slide_layouts[6])

    # 颜色 / 透明度
    @staticmethod
    def _rgb(c):
        from pptx.dml.color import RGBColor
        return RGBColor(c["r"], c["g"], c["b"])

    @staticmethod
    def _set_alpha(color_format_parent_xml, alpha: float):
        """在 <a:srgbClr> 下加 <a:alpha val=…/> (python-pptx 没有透明度 API)。"""
        from pptx.oxml.ns import qn
        if alpha >= 0.995:
            return
        for clr in color_format_parent_xml.iter(qn("a:srgbClr")):
            for old in clr.findall(qn("a:alpha")):
                clr.remove(old)
            el = clr.makeelement(qn("a:alpha"), {"val": str(int(max(0.0, alpha) * 100000))})
            clr.append(el)

    def fonts_for(self, fams: list[str]) -> tuple[str, str]:
        names = [f for f in fams or [] if f]
        low = [f.lower() for f in names]
        if self.keep_fonts and names:
            first = next((f for f in names if f.lower() not in ("sans-serif", "serif", "monospace")), self.font)
            return first, first
        if any(k in " ".join(low) for k in ("mono", "consolas", "courier")):
            return "Consolas", self.font
        serif = ("serif" in low and "sans-serif" not in low) or any(
            k in " ".join(low) for k in ("serif sc", "serif cjk", "song", "simsun", "mincho", "ming"))
        if serif:
            return "Georgia", "SimSun"
        latin = next((f for f in names if f.lower() in SAFE_LATIN), self.font)
        return latin, self.font

    # 图片来源 → 本地文件
    def local_image(self, src: str) -> Path | None:
        if not src:
            return None
        if src in self._dl:
            return self._dl[src]
        p = None
        try:
            if src.startswith("file://"):
                p = Path(urllib.parse.unquote(urllib.parse.urlparse(src).path))
            elif src.startswith("data:"):
                head, b64 = src.split(",", 1)
                ext = ".png" if "png" in head else ".svg" if "svg" in head else ".jpg"
                p = self.tmp / f"data{len(self._dl)}{ext}"
                p.write_bytes(base64.b64decode(b64) if ";base64" in head else urllib.parse.unquote(b64).encode())
            elif src.startswith(("http://", "https://")):
                p = self.tmp / f"dl{len(self._dl)}{Path(urllib.parse.urlparse(src).path).suffix or '.jpg'}"
                req = urllib.request.Request(src, headers={"User-Agent": "LiteCode-DocStudio/1.0"})
                with urllib.request.urlopen(req, timeout=30) as r:
                    p.write_bytes(r.read())
        except Exception as e:
            _ds.eprint(f"[html2pptx] 取图失败 {src[:80]}: {e}")
            p = None
        if p is not None and p.suffix.lower() in (".svg", ".webp", ".gif", ".avif") and p.exists():
            p = self._to_png(p)
        self._dl[src] = p if (p and p.exists()) else None
        return self._dl[src]

    def _to_png(self, p: Path) -> Path | None:
        try:
            from PIL import Image
            out = self.tmp / (p.stem + "_conv.png")
            Image.open(p).save(out)
            return out
        except Exception:
            return None                                # svg 等: 交给截图兜底 (调用方会跳过)

    # 元素
    def shape(self, slide, it: dict):
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.util import Emu
        if it.get("w", 0) < 0.5 or it.get("h", 0) < 0.5:
            return
        rad = it.get("radius") or 0
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if rad > 0.5 else MSO_SHAPE.RECTANGLE
        sp = slide.shapes.add_shape(kind, self.E(it["x"]), self.E(it["y"]), self.E(it["w"]), self.E(it["h"]))
        sp.shadow.inherit = False
        if rad > 0.5:
            sp.adjustments[0] = min(0.5, rad / max(1.0, min(it["w"], it["h"])))
        if it.get("fill"):
            sp.fill.solid()
            sp.fill.fore_color.rgb = self._rgb(it["fill"])
            self._set_alpha(sp.fill._xPr, it["fill"]["a"] * it.get("opacity", 1))
        else:
            sp.fill.background()
        if it.get("line"):
            sp.line.color.rgb = self._rgb(it["line"]["c"])
            sp.line.width = Emu(self.E(it["line"]["w"]))
            self._set_alpha(sp.line._ln, it["line"]["c"]["a"] * it.get("opacity", 1))
        else:
            sp.line.fill.background()
        sp.text_frame.text = ""

    def picture(self, slide, path: Path, it: dict):
        """按 object-fit / 裁切区域放图: 图片框 = 可见区域, crop_* 决定显示原图哪一部分。"""
        from PIL import Image
        nw, nh = it.get("natW") or 0, it.get("natH") or 0
        if not nw or not nh:
            with Image.open(path) as im:
                nw, nh = im.size
        x, y, w, h = it["x"], it["y"], it["w"], it["h"]
        fit = it.get("fit") or "fill"
        if fit == "fill":
            dx, dy, dw, dh = x, y, w, h
        else:
            if fit == "cover":
                s = max(w / nw, h / nh)
            elif fit in ("contain", "scale-down"):
                s = min(w / nw, h / nh)
                if fit == "scale-down":
                    s = min(1.0, s)
            else:                                          # none
                s = 1.0
            dw, dh = nw * s, nh * s
            px_, py_ = (it.get("pos") or [0.5, 0.5]) + [0.5] * (2 - len(it.get("pos") or []))
            dx, dy = x + (w - dw) * px_, y + (h - dh) * py_
        c = it.get("clip") or {"x": x, "y": y, "w": w, "h": h}
        vx0, vy0 = max(dx, x, c["x"]), max(dy, y, c["y"])
        vx1, vy1 = min(dx + dw, x + w, c["x"] + c["w"]), min(dy + dh, y + h, c["y"] + c["h"])
        if vx1 - vx0 < 1 or vy1 - vy0 < 1:
            return None
        pic = slide.shapes.add_picture(str(path), self.E(vx0), self.E(vy0), self.E(vx1 - vx0), self.E(vy1 - vy0))
        pic.crop_left, pic.crop_right = (vx0 - dx) / dw, (dx + dw - vx1) / dw
        pic.crop_top, pic.crop_bottom = (vy0 - dy) / dh, (dy + dh - vy1) / dh
        rad = it.get("radius") or 0
        if rad > 0.5:
            self._round_picture(pic, rad / max(1.0, min(vx1 - vx0, vy1 - vy0)))
        if it.get("opacity", 1) < 0.99:
            self._picture_alpha(pic, it["opacity"])
        return pic

    @staticmethod
    def _round_picture(pic, ratio: float):
        from pptx.oxml.ns import qn
        geom = pic._element.spPr.find(qn("a:prstGeom"))
        if geom is None:
            return
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = geom.makeelement(qn("a:avLst"), {})
            geom.append(av)
        for g in list(av):
            av.remove(g)
        av.append(av.makeelement(qn("a:gd"), {"name": "adj", "fmla": f"val {int(min(0.5, ratio) * 100000)}"}))

    @staticmethod
    def _picture_alpha(pic, alpha: float):
        from pptx.oxml.ns import qn
        blip = pic._element.find(".//" + qn("a:blip"))
        if blip is not None:
            blip.append(blip.makeelement(qn("a:alphaModFix"), {"amt": str(int(alpha * 100000))}))

    def _fill_runs(self, tf, runs: list[dict], align: str, line_h: float, opacity: float = 1.0):
        from pptx.enum.text import PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.util import Pt
        amap = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "end": PP_ALIGN.RIGHT,
                "justify": PP_ALIGN.JUSTIFY, "-webkit-center": PP_ALIGN.CENTER}
        p = tf.paragraphs[0]
        p.alignment = amap.get(align, PP_ALIGN.LEFT)
        if line_h:
            p.line_spacing = Pt(line_h * 0.75)
        for r in runs:
            if r.get("br"):
                p.add_line_break()
                continue
            run = p.add_run()
            run.text = r["text"]
            f = run.font
            f.size = Pt(max(1.0, r.get("size", 16) * 0.75))
            f.bold, f.italic = bool(r.get("bold")), bool(r.get("italic"))
            if r.get("underline"):
                f.underline = True
            latin, ea = self.fonts_for(r.get("fonts") or [])
            f.name = latin
            rpr = run._r.get_or_add_rPr()
            for tag, face in (("a:ea", ea), ("a:cs", latin)):
                el = rpr.find(qn(tag))
                if el is None:
                    el = rpr.makeelement(qn(tag), {})
                    rpr.append(el)
                el.set("typeface", face)
            if r.get("color"):
                f.color.rgb = self._rgb(r["color"])
                self._set_alpha(rpr, r["color"]["a"] * opacity)
            if r.get("spacing"):
                rpr.set("spc", str(int(r["spacing"] * 0.75 * 100)))
            if r.get("sup") or r.get("sub"):
                rpr.set("baseline", "30000" if r.get("sup") else "-25000")
            if r.get("href", "").startswith(("http://", "https://")):
                run.hyperlink.address = r["href"]

    def text(self, slide, it: dict):
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
        single = it.get("lines", 1) <= 1
        x, w = it["x"], it["w"]
        if not single:                                  # Office 字宽略不同, 多行留一点余量防多折一行
            extra = max(4.0, w * 0.03)
            if it.get("align") in ("center", "-webkit-center"):
                x -= extra / 2
            elif it.get("align") in ("right", "end"):
                x -= extra
            w += extra
        tb = slide.shapes.add_textbox(self.E(x), self.E(it["y"]), self.E(max(w, 2)), self.E(max(it["h"], 2)))
        tf = tb.text_frame
        tf.word_wrap = not single
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        self._fill_runs(tf, it["runs"], it.get("align", "left"), it.get("lineHeight", 0), it.get("opacity", 1))

    def table(self, slide, it: dict):
        from pptx.util import Emu
        rows = [r for r in it.get("rows") or [] if r.get("cells")]
        if not rows:
            return
        edges = sorted({round(c["x"]) for r in rows for c in r["cells"]} |
                       {round(c["x"] + c["w"]) for r in rows for c in r["cells"]})
        cols = []                                          # 合并 2px 内的近邻边
        for e in edges:
            if not cols or e - cols[-1] > 2:
                cols.append(e)
        if len(cols) < 2:
            return
        x0, y0 = cols[0], rows[0]["y"]
        gf = slide.shapes.add_table(len(rows), len(cols) - 1, self.E(x0), self.E(y0),
                                    self.E(cols[-1] - x0), self.E(sum(r["h"] for r in rows)))
        tbl = gf.table
        tbl.first_row = tbl.horz_banding = False
        for j in range(len(cols) - 1):
            tbl.columns[j].width = Emu(self.E(cols[j + 1] - cols[j]))
        for i, r in enumerate(rows):
            tbl.rows[i].height = Emu(self.E(r["h"]))
        near = lambda v: min(range(len(cols)), key=lambda k: abs(cols[k] - v))  # noqa: E731
        for i, r in enumerate(rows):
            for c in r["cells"]:
                j0, j1 = near(c["x"]), near(c["x"] + c["w"]) - 1
                i1 = min(len(rows) - 1, i + max(1, c.get("rowSpan", 1)) - 1)
                if j1 < j0:
                    continue
                cell = tbl.cell(i, j0)
                if (j1 > j0 or i1 > i) and not cell.is_spanned:
                    try:
                        cell.merge(tbl.cell(i1, j1))
                    except Exception:
                        pass
                self._cell(cell, c)

    def _cell(self, cell, c: dict):
        from pptx.enum.text import MSO_ANCHOR
        from pptx.oxml.ns import qn
        from pptx.util import Emu
        if c.get("fill"):
            cell.fill.solid()
            cell.fill.fore_color.rgb = self._rgb(c["fill"])
        else:
            cell.fill.background()
        pt, pr, pb, pl = c.get("pad") or [4, 6, 4, 6]
        cell.margin_top, cell.margin_right = Emu(self.E(pt)), Emu(self.E(pr))
        cell.margin_bottom, cell.margin_left = Emu(self.E(pb)), Emu(self.E(pl))
        cell.vertical_anchor = {"top": MSO_ANCHOR.TOP, "bottom": MSO_ANCHOR.BOTTOM}.get(c.get("valign"), MSO_ANCHOR.MIDDLE)
        tf = cell.text_frame
        tf.word_wrap = True
        size = next((r.get("size") for r in c.get("runs") or [] if r.get("size")), 14)
        self._fill_runs(tf, c.get("runs") or [], c.get("align", "left"), size * 1.3)
        tcPr = cell._tc.get_or_add_tcPr()
        for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
            for old in tcPr.findall(qn(tag)):
                tcPr.remove(old)
        for k, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
            b = c.get("border") if tag == "a:lnB" else None
            ln = tcPr.makeelement(qn(tag), {"w": str(self.E(b["w"])) if b else "0"})
            if b and b.get("c"):
                sf = ln.makeelement(qn("a:solidFill"), {})
                col = sf.makeelement(qn("a:srgbClr"), {"val": "%02X%02X%02X" % (b["c"]["r"], b["c"]["g"], b["c"]["b"])})
                sf.append(col)
                ln.append(sf)
            else:
                ln.append(ln.makeelement(qn("a:noFill"), {}))
            tcPr.insert(k, ln)

    def background(self, slide, fill: dict | None):
        if fill:
            bg = slide.background.fill
            bg.solid()
            bg.fore_color.rgb = self._rgb(fill)

    def hidden_text(self, slide, text: str):
        """高保真页: 文字框放在最底层 (被整页图盖住), 不可见但能搜索/复制。"""
        from pptx.util import Pt
        if not text.strip():
            return
        tb = slide.shapes.add_textbox(0, 0, self.prs.slide_width, self.prs.slide_height)
        tb.name = "hidden-text"
        tf = tb.text_frame
        tf.word_wrap = True
        tf.text = text
        for p in tf.paragraphs:
            for r in p.runs:
                r.font.size = Pt(10)

    def notes(self, slide, text: str):
        if text.strip():
            slide.notes_slide.notes_text_frame.text = text.strip()


def convert(html: str, out: str, mode: str = "auto", font: str = "Microsoft YaHei",
            keep_fonts: bool = False, scale: float = 2.0, selector: str | None = None) -> dict:
    tmp = Path(tempfile.mkdtemp(prefix="html2pptx_"))
    report = {"pptx": str(Path(out).resolve()), "modes": [], "reasons": []}
    try:
        kind = _ds.sniff_kind(html)
        vw, vh = _ds.page_size_for(kind)
        with _ds.open_page(html, width=vw, height=vh, scale=scale) as (page, _b):
            sel = selector or ("section.slide" if page.evaluate("!!document.querySelector('section.slide')")
                               else "section.page")
            slides = extract(page, sel)
            if not slides:
                raise SystemExit(f"没有找到 {sel}: 幻灯片请每页用 <section class=\"slide\">")
            page.add_style_tag(content=CAP_CSS)
            plans = []
            for i, s in enumerate(slides):
                m, why = choose_mode(s, mode)
                shots = {}
                if m == "image":
                    shots["full"] = shoot_section(page, sel, i, tmp / f"s{i}_full.png")
                else:
                    for it in s["items"]:
                        if it["t"] == "cap":
                            shots[it["id"]] = capture_el(page, it["id"], it["mode"], tmp / f"cap{it['id']}.png")
                plans.append((s, m, why, shots))
        b = Builder(slides[0]["w"], slides[0]["h"], font, keep_fonts, tmp)
        for s, m, why, shots in plans:
            slide = b.new_slide()
            sx = (slides[0]["w"] / s["w"]) if s.get("w") else 1.0
            if m == "image":
                b.hidden_text(slide, s.get("text", ""))
                slide.shapes.add_picture(str(shots["full"]), 0, 0, b.prs.slide_width, b.prs.slide_height)
                b.notes(slide, "\n\n".join(x for x in (s.get("notes", ""), s.get("text", "")) if x))
            else:
                for it in s["items"]:
                    if sx != 1.0:
                        for k in ("x", "y", "w", "h"):
                            it[k] = it[k] * sx
                    t = it["t"]
                    if t == "shape" and it.get("root") and not it.get("line") and (it.get("radius") or 0) < 1:
                        b.background(slide, it.get("fill"))
                    elif t == "shape":
                        b.shape(slide, it)
                    elif t == "image":
                        p = b.local_image(it.get("src", ""))
                        if p:
                            b.picture(slide, p, it)
                    elif t == "cap" and shots.get(it["id"]):
                        it2 = dict(it, fit="fill", natW=0, natH=0)
                        b.picture(slide, shots[it["id"]], it2)
                    elif t == "text":
                        b.text(slide, it)
                    elif t == "table":
                        b.table(slide, it)
                b.notes(slide, s.get("notes", ""))
            report["modes"].append(m)
            report["reasons"].append(why)
        Path(out).parent.mkdir(parents=True, exist_ok=True)
        b.prs.save(out)
        report["slides"] = len(plans)
        return report
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def main(argv=None):
    ap = argparse.ArgumentParser(description="doc-studio: HTML 幻灯片 → PPTX")
    ap.add_argument("html")
    ap.add_argument("-o", "--out", required=True)
    ap.add_argument("--mode", default="auto", choices=["auto", "editable", "image"])
    ap.add_argument("--font", default="Microsoft YaHei", help="中文字体 (PPT 里显示用), 默认微软雅黑")
    ap.add_argument("--keep-fonts", action="store_true", help="保留 HTML 里写的字体名, 不映射")
    ap.add_argument("--scale", type=float, default=2.0, help="截图倍率")
    ap.add_argument("--selector", help="每页的选择器 (默认 section.slide, 没有则 section.page)")
    a = ap.parse_args(argv)
    print(_ds.dump_json(convert(a.html, a.out, a.mode, a.font, a.keep_fonts, a.scale, a.selector)))


if __name__ == "__main__":
    main()
