"""plugins/kb/text.py — 文档解析 + 切分。

解析结果是段列表 Segment(text, unit, num) —— 页 / 段 / 幻灯片 / 工作表, 拼成全文时记下每段的起止偏移,
切块后按块的偏移反查出处 ("第3页" / "第12-14段")。全部本地解析, 不依赖外部 OCR/MinIO 服务:
  pdf  pypdf 逐页取文字; 文字过少的页 (扫描件) 用 pypdfium2 渲染后走视觉识别
  docx python-docx (段落 + 表格)      pptx python-pptx (每页文本框)    xlsx openpyxl (每个工作表)
  txt/md/csv/json/log/代码  直接读 (utf-8 → gb18030 兜底)            html  去标签
  图片  走当前模型的视觉识别 (core/tools/vision.py 的 vision_ocr, 配置 kb.ocr=false 可关)

切分 (沿用原知识库项目的句末断句思路, 参数改小): 默认 450 字一块, 相邻块重叠 80 字,
优先在句末标点处断开, 其次逗号/空白; 尾巴不足 1/4 块并入前一块。
"""
from __future__ import annotations

import asyncio
import html as _html
import re
import tempfile
import unicodedata
from dataclasses import dataclass
from pathlib import Path

TEXT_EXT = {".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".log", ".yaml", ".yml", ".ini", ".py",
            ".js", ".ts", ".java", ".go", ".c", ".cpp", ".h", ".sql", ".sh", ".xml", ".rst"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"}
SUPPORTED = TEXT_EXT | IMAGE_EXT | {".pdf", ".docx", ".pptx", ".xlsx", ".html", ".htm"}

_OCR_PROMPT = "请逐字输出图中所有文字, 保持原版式 (标题/段落/表格/列表结构), 不要翻译, 不要解释。"
_MIN_PAGE_TEXT = 20     # 一页 PDF 抽出的文字少于此数 → 当扫描页做 OCR


@dataclass
class Segment:
    text: str
    unit: str = ""      # 页 / 段 / 张幻灯片 / "" (工作表等用 label)
    num: int = 0
    label: str = ""


# ── OCR ───────────────────────────────────────────────────────────

def _ocr(path: str) -> str:
    """同步调用视觉识别 (在入库线程里跑, 自建事件循环)。失败返回空串并带错误前缀给调用方判断。"""
    try:
        try:
            from tools.vision import _do_vision_ocr        # web_ui / server 进程: core/ 在 sys.path
        except ImportError:
            from core.tools.vision import _do_vision_ocr
    except Exception as e:  # noqa: BLE001
        return f"ERROR: 视觉识别模块不可用: {e}"
    try:
        return asyncio.run(_do_vision_ocr(path=path, prompt=_OCR_PROMPT, max_tokens=4096)) or ""
    except Exception as e:  # noqa: BLE001
        return f"ERROR: 视觉识别失败: {e}"


def _ocr_ok(s: str) -> str:
    if s.startswith("ERROR"):
        raise RuntimeError(s)
    return s.strip()


# ── 解析 ──────────────────────────────────────────────────────────

def _read_text(p: Path) -> str:
    raw = p.read_bytes()
    for enc in ("utf-8-sig", "gb18030"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", "replace")


def _paragraphs(text: str) -> list[Segment]:
    parts = [x.strip() for x in re.split(r"\n\s*\n", text.replace("\r\n", "\n"))]
    return [Segment(x, "段", i) for i, x in enumerate((x for x in parts if x), 1)]


def _parse_pdf(p: Path, ocr: bool) -> list[Segment]:
    from pypdf import PdfReader
    reader = PdfReader(str(p))
    segs, scanned = [], []
    for i, page in enumerate(reader.pages, 1):
        try:
            t = (page.extract_text() or "").strip()
        except Exception:  # noqa: BLE001  个别页解析失败不拖垮整份
            t = ""
        if len(re.sub(r"\s", "", t)) < _MIN_PAGE_TEXT:
            scanned.append(i)
        segs.append(Segment(t, "页", i))
    if scanned and ocr:
        import pypdfium2 as pdfium
        doc = pdfium.PdfDocument(str(p))
        with tempfile.TemporaryDirectory() as td:
            for i in scanned:
                img = doc[i - 1].render(scale=2).to_pil()
                fp = Path(td) / f"p{i}.png"
                img.save(fp)
                text = _ocr_ok(_ocr(str(fp)))
                if text:
                    segs[i - 1] = Segment(text, "页", i)
    return segs


def _parse_docx(p: Path) -> list[Segment]:
    import docx
    d = docx.Document(str(p))
    texts = [x.text.strip() for x in d.paragraphs]
    for tb in d.tables:
        for row in tb.rows:
            cells = []
            for c in row.cells:
                v = c.text.strip()
                if v and (not cells or cells[-1] != v):     # 合并单元格会重复出现
                    cells.append(v)
            if cells:
                texts.append(" | ".join(cells))
    return [Segment(t, "段", i) for i, t in enumerate((t for t in texts if t), 1)]


def _parse_pptx(p: Path) -> list[Segment]:
    from pptx import Presentation
    segs = []
    for i, slide in enumerate(Presentation(str(p)).slides, 1):
        lines = [sh.text_frame.text.strip() for sh in slide.shapes if getattr(sh, "has_text_frame", False)]
        segs.append(Segment("\n".join(x for x in lines if x), "张幻灯片", i))
    return segs


def _parse_xlsx(p: Path) -> list[Segment]:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    segs = []
    for ws in wb.worksheets:
        rows = []
        for r in ws.iter_rows(values_only=True):
            cells = [str(v).strip() for v in r if v is not None and str(v).strip()]
            if cells:
                rows.append(" | ".join(cells))
            if len(rows) >= 20000:
                break
        segs.append(Segment("\n".join(rows), label=f"工作表「{ws.title}」"))
    wb.close()
    return segs


def _parse_html(p: Path) -> list[Segment]:
    s = _read_text(p)
    s = re.sub(r"(?is)<(script|style|noscript)[^>]*>.*?</\1>", " ", s)
    s = re.sub(r"(?i)<br\s*/?>|</(p|div|h[1-6]|li|tr|section|article)>", "\n\n", s)
    s = _html.unescape(re.sub(r"<[^>]+>", " ", s))
    s = re.sub(r"[ \t　]+", " ", s)
    return _paragraphs(s)


def parse(path: str | Path, ocr: bool = True) -> list[Segment]:
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        return _parse_pdf(p, ocr)
    if ext == ".docx":
        return _parse_docx(p)
    if ext == ".pptx":
        return _parse_pptx(p)
    if ext == ".xlsx":
        return _parse_xlsx(p)
    if ext in (".html", ".htm"):
        return _parse_html(p)
    if ext in IMAGE_EXT:
        if not ocr:
            raise RuntimeError("图片需要视觉识别, 但 kb.ocr 已关闭")
        return [Segment(_ocr_ok(_ocr(str(p))), label="图片识别")]
    if ext in TEXT_EXT:
        return _paragraphs(_read_text(p))
    raise RuntimeError(f"不支持的文件类型 {ext or '(无扩展名)'}; 支持: {' '.join(sorted(SUPPORTED))}")


# Chrome/Office 导出的 PDF 常把部分汉字的 ToUnicode 映射成部首字符 (门→⻔ 青→⻘ 人→⼈), 抽出来的文字
# 检索不到。康熙部首 (U+2F00-2FD5) NFKC 能还原; 部首补充区 (U+2E80-2EFF) 没有兼容映射, 常见的手工对照。
_RADICALS = str.maketrans("⺠⻅⻆⻉⻋⻊⻓⻔⻘⻚⻛⻜⻝⻢⻥⻦⻧⻨⻩⻪⻬⻮⻰⺼⺗",
                          "民见角贝车足长门青页风飞食马鱼鸟卤麦黄黾齐齿龙月心")


def clean(t: str) -> str:
    """NFKC (全角转半角、康熙部首转常用字) + 部首补充区对照 + 去掉控制字符。"""
    t = unicodedata.normalize("NFKC", t or "").translate(_RADICALS)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f�]", "", t)


def join(segs: list[Segment]) -> tuple[str, list[tuple[int, int, Segment]]]:
    """段 → 全文 + [(起, 止, 段)]; 段间用空行隔开。"""
    parts, spans, pos = [], [], 0
    for s in segs:
        t = clean(s.text).strip()
        if not t:
            continue
        if parts:
            parts.append("\n\n")
            pos += 2
        parts.append(t)
        spans.append((pos, pos + len(t), s))
        pos += len(t)
    return "".join(parts), spans


def locate(spans, start: int, end: int) -> tuple[str, int | None]:
    """块 [start, end) → (出处文字, 页码或 None)。"""
    hit = [s for a, b, s in spans if a < end and b > start] or [s for a, b, s in spans if a <= start][-1:]
    if not hit:
        return "", None
    first, last = hit[0], hit[-1]
    page = first.num if first.unit == "页" else None
    if not first.unit:
        return first.label, None
    if last.unit == first.unit and last.num != first.num:
        return f"第{first.num}-{last.num}{first.unit}", page
    return f"第{first.num}{first.unit}", page


# ── 切分 ──────────────────────────────────────────────────────────

_SENT_END = set("。！？!?；;…\n")
_SOFT_END = set("，,、：: \t")


def split(text: str, size: int = 450, overlap: int = 80) -> list[tuple[int, int, str]]:
    """→ [(起, 止, 块文本)]。起止是全文偏移 (块文本是 text[起:止] 去首尾空白)。"""
    size = max(50, int(size))
    overlap = max(0, min(int(overlap), size // 2))
    n, out, start = len(text), [], 0
    while start < n:
        while start < n and text[start].isspace():
            start += 1
        if start >= n:
            break
        end = min(start + size, n)
        if n - end < size // 4:                      # 尾巴太短: 并入本块
            end = n
        else:
            lo = start + int(size * 0.6)
            cut = next((p + 1 for p in range(end - 1, lo - 1, -1)
                        if text[p] in _SENT_END or (text[p] == "." and (p + 1 >= n or text[p + 1].isspace()))), -1)
            if cut < 0:
                cut = next((p + 1 for p in range(end - 1, lo - 1, -1) if text[p] in _SOFT_END), -1)
            if cut > 0:
                end = cut
        chunk = text[start:end].strip()
        if chunk:
            out.append((start, end, chunk))
        if end >= n:
            break
        nxt = end - overlap
        for p in range(nxt, end):                    # 重叠区里尽量从句首开始
            if p > 0 and text[p - 1] in _SENT_END:
                nxt = p
                break
        start = max(nxt, start + 1)
    return out


# ── 关键词 (中文 bigram + 英文/数字词), 入库和查询共用 ─────────────────

_WORD = re.compile(r"[a-z0-9][a-z0-9_.\-]*[a-z0-9]|[a-z0-9]")
_CJK_RUN = re.compile(r"[㐀-䶿一-鿿豈-﫿]+")
_STOP = set("的了是在和与及或也就都而且之其这那个为把被让给对从向于以上下中有无不吗呢吧啊么")
_GLUE = set("的了是和与及或吗呢吧啊么")          # 含这些字的 bigram 基本是虚词搭配 (是什/的年), 不当关键词
_QUESTION = {"什么", "怎么", "怎样", "如何", "多少", "哪些", "哪个", "哪里", "是否", "能否", "可以", "请问", "一下"}


def terms(text: str, query: bool = False) -> list[str]:
    """中文 bigram + 英文/数字词。query=True 时再去掉疑问词 (什么/多少/如何…), 只留内容词。"""
    t = (text or "").lower()
    out = _WORD.findall(t)
    for run in _CJK_RUN.findall(t):
        if len(run) == 1:
            if run not in _STOP:
                out.append(run)
            continue
        out.extend(b for b in (run[i:i + 2] for i in range(len(run) - 1))
                   if not (b[0] in _GLUE or b[1] in _GLUE) and not (query and b in _QUESTION))
    return out
